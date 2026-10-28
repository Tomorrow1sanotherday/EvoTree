"""Deterministic held-out parameter/layout/context variants, never seed data.

These are custom tasks using OSWorld environments and objective file metrics.
They are not an official benchmark split or unseen-template evaluation.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
from pathlib import Path
import shutil
import uuid

from experiments.m2_small_eval.diagnostic_pilot import write_json


CALC = [
    dict(level="parameters", label="Operating Cost", column=1, count=12, sheet="Data", target="CostSummary", start="Data"),
    dict(level="layout", label="Units Sold", column=4, count=17, sheet="Transactions", target="UnitsReport", start="Transactions"),
    dict(level="context", label="Tax Amount", column=2, count=9, sheet="Quarterly", target="TaxReview", start="Scratch"),
]
FONT = [
    dict(level="parameters", slide=14, sizes=[44, 22], boxes=[6, 7]),
    dict(level="layout", slide=5, sizes=[36, 18], boxes=[2, 3]),
    dict(level="context", slide=9, sizes=[48, 24], boxes=[1, 2]),
]


def save_config(case_id, filename, instruction, domain, metric, options, base_url):
    guest = "/home/user/" + filename
    return dict(id=case_id, snapshot=domain, instruction=instruction,
        source="custom held-out variation; not official OSWorld result",
        config=[{"type": "download", "parameters": {"files": [{"url": f"{base_url}/{case_id}/input/{filename}", "path": guest}]}},
                {"type": "open", "parameters": {"path": guest}}],
        related_apps=[domain], proxy=False, fixed_ip=False,
        evaluator={"postconfig": [
            {"type": "activate_window", "parameters": {"window_name": filename + " - LibreOffice " + ("Calc" if domain.endswith("calc") else "Impress"), "strict": True}},
            {"type": "execute", "parameters": {"command": ["python", "-c", "import pyautogui; pyautogui.hotkey('ctrl','s')"]}},
            {"type": "sleep", "parameters": {"seconds": 0.5}}],
            "func": metric, "result": {"type": "vm_file", "path": guest, "dest": filename},
            "expected": {"type": "cloud_file", "path": f"{base_url}/{case_id}/gold/{filename}", "dest": "gold-" + filename},
            "options": options})


def make_cases(output, original_pptx, base_url, *, calc_variants=None, font_variants=None):
    import openpyxl
    from openpyxl.utils import get_column_letter
    from pptx import Presentation
    from pptx.util import Inches, Pt
    output.mkdir(parents=True, exist_ok=False)
    cases = []
    for index, values in enumerate(CALC if calc_variants is None else calc_variants):
        case_id = values.get("id", f"transfer-calc-{values['level']}")
        filename = f"TransferLedger{index+1}.xlsx"
        folder = output / case_id
        (folder / "input").mkdir(parents=True)
        (folder / "gold").mkdir()
        book = openpyxl.Workbook()
        source = book.active
        source.title = values["sheet"]
        for col in range(1, 5):
            source.cell(1, col, values["label"] if col == values["column"] else f"Control {col}")
            for row in range(2, values["count"]+1):
                source.cell(row, col, (index+3)*1000 + row*37 + col*11)
        if values["level"] == "context":
            blank = book.create_sheet("Scratch")
            blank["A1"] = "Unrelated notes; data are in Quarterly"
            book.active = 1
        book.save(folder / "input" / filename)
        original_sheets = book.sheetnames[:]
        target = book.create_sheet(values["target"])
        for row in range(1, values["count"]+1):
            target.cell(row, 1, source.cell(row, values["column"]).value)
        book.save(folder / "gold" / filename)
        rules = [{"type": "sheet_data", "sheet_idx0": "RN"+name, "sheet_idx1": "EN"+name}
                 for name in original_sheets + [values["target"]]]
        instruction = f'Copy the "{values["label"]}" column along with the header to a new sheet named "{values["target"]}".'
        config = save_config(case_id, filename, instruction, "libreoffice_calc", "compare_table", {"rules": rules}, base_url)
        # Correct command sequence is a harness fixture only. It never enters
        # memory or model input and is not used to select promising tasks.
        calls = [
            {"action_type": "libreoffice_calc.switch_active_sheet", "parameters": {"sheet_name": values["target"]}},
            {"action_type": "libreoffice_calc.switch_active_sheet", "parameters": {"sheet_name": values["sheet"]}},
            {"action_type": "libreoffice_calc2.copy_cells_between_sheets", "parameters": {
                "source_range": f'{get_column_letter(values["column"])}1:{get_column_letter(values["column"])}{values["count"]}',
                "target_sheet_name": values["target"], "target_start_cell": "A1"}},
            {"action_type": "libreoffice_calc.save", "parameters": {}}]
        cases.append(dict(id=case_id, family="calc", level=values["level"], config=config,
            smoke_calls=calls, source_variant=values, input=str(folder/"input"/filename), gold=str(folder/"gold"/filename)))
    for index, values in enumerate(FONT if font_variants is None else font_variants):
        case_id = values.get("id", f"transfer-font-{values['level']}")
        filename = f"TransferSlides{index+1}.pptx"
        folder = output / case_id
        (folder/"input").mkdir(parents=True)
        (folder/"gold").mkdir()
        if values["level"] == "parameters" and values.get("use_original", True):
            presentation = Presentation(str(original_pptx))
        else:
            presentation = Presentation()
            for number in range(1, 17):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                offset = values["boxes"][0] if number == values["slide"] else 1
                # Decorative blank shapes shift the actual textbox indices.
                for decorative in range(offset):
                    shape = slide.shapes.add_shape(1, Inches(0.2+decorative*0.2), Inches(0.2), Inches(0.1), Inches(0.1))
                    shape.fill.background()
                for j in range(2):
                    shape = slide.shapes.add_textbox(Inches(1), Inches(1.5+j*2), Inches(7), Inches(1.2))
                    shape.text = (values.get('text_prefix', '') +
                                  (f"Quarter {number} overview" if j == 0 else f"Planning notes for section {number}"))
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(30 if j == 0 else 16)
        slide = presentation.slides[values["slide"]-1]
        texts = [slide.shapes[box].text for box in values["boxes"]]
        for j, box in enumerate(values["boxes"]):
            for paragraph in slide.shapes[box].text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(values["sizes"][j] if values["level"] == "context" and j == 0 else 30-j*14)
        presentation.save(folder/"input"/filename)
        for j, box in enumerate(values["boxes"]):
            for paragraph in slide.shapes[box].text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(values["sizes"][j])
        presentation.save(folder/"gold"/filename)
        instruction = f'Make the first textbox font size {values["sizes"][0]} pt while the second {values["sizes"][1]} pt on slide {values["slide"]}.'
        options = dict(examine_shape=False, examine_font_size=True, examine_text=True,
            examine_font_name=False, examine_color_rgb=False, examine_indent=False,
            examine_font_bold=False, examine_font_italic=False, examine_font_underline=False,
            examine_strike_through=False, examine_bullets=False, examine_background_color=False,
            examine_note=False, examine_run_count=False, examine_alignment=False, enable_debug=False)
        config = save_config(case_id, filename, instruction, "libreoffice_impress", "compare_pptx_files", options, base_url)
        calls = [{"action_type": "libreoffice_impress.write_text", "parameters": {
            "content": texts[j], "page_index": values["slide"], "box_index": box,
            "size": values["sizes"][j], "append": False}} for j, box in enumerate(values["boxes"])]
        calls.append({"action_type": "libreoffice_impress.save", "parameters": {}})
        cases.append(dict(id=case_id, family="font", level=values["level"], config=config,
            smoke_calls=calls, source_variant=values, input=str(folder/"input"/filename), gold=str(folder/"gold"/filename)))
    for case in cases:
        write_json(output/case["id"]/"task.json",case["config"])
        case["input_sha256"] = hashlib.sha256(Path(case["input"]).read_bytes()).hexdigest()
        case["gold_sha256"] = hashlib.sha256(Path(case["gold"]).read_bytes()).hexdigest()
    write_json(output/"cases-private.json", cases)
    write_json(output/"split.json", dict(setting="custom same-template held-out variants", train_source="fixed historical seed episodes only",
        test=[dict(id=c["id"],family=c["family"],level=c["level"],instruction=c["config"]["instruction"],input_sha256=c["input_sha256"]) for c in cases],
        prompt_access="task instruction and live observation only; no gold or fixture calls", official_benchmark=False))
    return cases


def prepare_cache(case, destination):
    """Local immutable inputs and expected file; never a scored result."""
    config = case["config"]
    folder = destination/case["id"]
    folder.mkdir(parents=True,exist_ok=True)
    file = config["config"][0]["parameters"]["files"][0]
    name = f"{uuid.uuid5(uuid.NAMESPACE_URL,file['url'])}_{Path(file['path']).name}"
    shutil.copy2(case["input"],folder/name)
    shutil.copy2(case["gold"],folder/config["evaluator"]["expected"]["dest"])


if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",type=Path,required=True)
    parser.add_argument("--original-pptx",type=Path,required=True)
    parser.add_argument("--base-url",default="http://20.20.20.1:18766")
    args=parser.parse_args()
    print(json.dumps(dict(cases=len(make_cases(args.output,args.original_pptx,args.base_url)))))
